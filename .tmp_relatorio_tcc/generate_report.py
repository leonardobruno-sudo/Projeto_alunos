from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = Path(__file__).resolve().parent / "assets"
OUT_FILE = ROOT / "entregaveis" / "Relatorio_Tecnico_SAC_IFRR.docx"

BLUE = "#294E82"
MID_BLUE = "#53759F"
LIGHT_BLUE = "#EAF1F8"
INK = "#111111"
GRAY = "#666666"
LIGHT_GRAY = "#F2F2F2"
FONT = "Times New Roman"

SECTION_TITLES = [
    "1 INTRODUÇÃO",
    "2 MÉTODO DE DESENVOLVIMENTO E DELIMITAÇÃO",
    "3 VISÃO GERAL DO SISTEMA",
    "4 ENGENHARIA E ALTERAÇÕES IMPLEMENTADAS",
    "5 MODELAGEM E DIAGRAMAS",
    "6 SEGURANÇA, AUTENTICAÇÃO E CONTROLE DE ACESSO",
    "7 VALIDAÇÃO E QUALIDADE",
    "8 LIMITAÇÕES E OPORTUNIDADES DE MELHORIA",
    "9 CONSIDERAÇÕES FINAIS",
    "REFERÊNCIAS",
    "APÊNDICE A - MATRIZ DE FUNCIONALIDADES",
    "APÊNDICE B - ROTEIRO DE OPERAÇÃO E VERIFICAÇÃO",
]

FIGURES = [
    ("Figura 1", "Diagrama de casos de uso do Sistema de Alunos Cotistas (SAC)"),
    ("Figura 2", "Diagrama lógico de relacionamentos de dados"),
    ("Figura 3", "Arquitetura em camadas do SAC"),
]


def font_path(bold: bool = False) -> str:
    candidates = [
        r"C:\\Windows\\Fonts\\timesbd.ttf" if bold else r"C:\\Windows\\Fonts\\times.ttf",
        r"C:\\Windows\\Fonts\\arialbd.ttf" if bold else r"C:\\Windows\\Fonts\\arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return "arial.ttf"


def pfont(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path(bold), size=size)


def pil_color(value: str) -> str:
    return value if value.startswith("#") else f"#{value}"


def text_box(draw: ImageDraw.ImageDraw, box, text, *, fill="FFFFFF", outline=BLUE,
             title=False, radius=20, padding=18, align="center", font_size=None):
    x0, y0, x1, y1 = box
    draw.rounded_rectangle(box, radius=radius, fill=pil_color(fill), outline=pil_color(outline), width=3)
    size = font_size or (30 if title else 24)
    regular = pfont(size, False)
    bold = pfont(size, True)
    lines = text.split("\n")
    line_heights = []
    total = 0
    for index, line in enumerate(lines):
        font = bold if index == 0 and title else regular
        bbox = draw.textbbox((0, 0), line, font=font)
        height = bbox[3] - bbox[1]
        line_heights.append((font, height))
        total += height + 8
    y = y0 + max(padding, (y1 - y0 - total) / 2)
    for index, line in enumerate(lines):
        font, height = line_heights[index]
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        if align == "left":
            x = x0 + padding
        else:
            x = (x0 + x1 - width) / 2
        draw.text((x, y), line, font=font, fill=INK)
        y += height + 8


def ellipse(draw: ImageDraw.ImageDraw, box, text, *, fill="FFFFFF", outline=BLUE):
    x0, y0, x1, y1 = box
    draw.ellipse(box, fill=pil_color(fill), outline=pil_color(outline), width=3)
    font = pfont(21, False)
    lines = text.split("\n")
    total = sum(draw.textbbox((0, 0), line, font=font)[3] - draw.textbbox((0, 0), line, font=font)[1] + 5 for line in lines)
    y = (y0 + y1 - total) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        height = bbox[3] - bbox[1]
        draw.text(((x0 + x1 - width) / 2, y), line, font=font, fill=INK)
        y += height + 5


def actor(draw: ImageDraw.ImageDraw, x: int, y: int, label: str):
    color = BLUE
    draw.ellipse((x - 20, y, x + 20, y + 40), outline=color, width=3)
    draw.line((x, y + 40, x, y + 105), fill=color, width=3)
    draw.line((x - 38, y + 62, x + 38, y + 62), fill=color, width=3)
    draw.line((x, y + 105, x - 35, y + 145), fill=color, width=3)
    draw.line((x, y + 105, x + 35, y + 145), fill=color, width=3)
    font = pfont(22, True)
    bbox = draw.textbbox((0, 0), label, font=font)
    draw.text((x - (bbox[2] - bbox[0]) / 2, y + 158), label, font=font, fill=INK)


def arrow(draw: ImageDraw.ImageDraw, start, end, *, color=GRAY, width=2, dashed=False):
    x0, y0 = start
    x1, y1 = end
    if dashed:
        steps = 12
        for index in range(steps):
            if index % 2 == 0:
                t0 = index / steps
                t1 = (index + 1) / steps
                draw.line((x0 + (x1 - x0) * t0, y0 + (y1 - y0) * t0,
                           x0 + (x1 - x0) * t1, y0 + (y1 - y0) * t1), fill=color, width=width)
    else:
        draw.line((x0, y0, x1, y1), fill=color, width=width)
    # Simple triangular arrowhead.
    dx = x1 - x0
    dy = y1 - y0
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    point1 = (x1, y1)
    point2 = (x1 - 15 * ux + 7 * px, y1 - 15 * uy + 7 * py)
    point3 = (x1 - 15 * ux - 7 * px, y1 - 15 * uy - 7 * py)
    draw.polygon((point1, point2, point3), fill=color)


def save_image(image: Image.Image, filename: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output = ASSET_DIR / filename
    image.save(output, "PNG", optimize=True)
    return output


def make_use_case_diagram() -> Path:
    image = Image.new("RGB", (1900, 1250), "WHITE")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((310, 80, 1590, 1150), radius=30, outline=BLUE, width=4)
    title_font = pfont(30, True)
    draw.text((790, 105), "Sistema de Alunos Cotistas (SAC)", font=title_font, fill=BLUE)

    # Actors
    actor(draw, 130, 360, "Aluno")
    actor(draw, 130, 760, "Professor")
    actor(draw, 1760, 360, "Diretor(a)")
    actor(draw, 1760, 760, "Administrador(a)")

    cases = {
        "login": (590, 205, 930, 325, "Autenticar-se\ne encerrar sessão"),
        "password": (970, 205, 1310, 325, "Alterar\nsenha"),
        "self": (420, 450, 800, 585, "Consultar dados e\ndesempenho próprio"),
        "stats": (1010, 450, 1420, 585, "Consultar estatísticas\nconforme o escopo"),
        "student": (430, 690, 830, 825, "Gerir alunos e dados\nacadêmicos no escopo"),
        "import": (1010, 690, 1420, 825, "Cadastrar e importar\nalunos por CSV"),
        "history": (470, 915, 865, 1045, "Salvar histórico\npor período"),
        "admin": (1000, 900, 1445, 1055, "Administrar equipe,\nbackup e permissões"),
    }
    for key, (x0, y0, x1, y1, label) in cases.items():
        ellipse(draw, (x0, y0, x1, y1), label, fill=LIGHT_BLUE if key in {"stats", "history", "admin"} else "FFFFFF")

    # Associations. Lines intentionally avoid arrowheads because UML associations are undirected.
    def assoc(a, b):
        draw.line((*a, *b), fill=GRAY, width=2)

    # Only profile-specific links are drawn so the diagram remains legible. The
    # common operations are documented in the note below the ellipses.
    assoc((170, 440), (420, 515))
    assoc((170, 840), (430, 757))
    assoc((1730, 440), (1420, 517))
    assoc((1730, 840), (1445, 980))
    assoc((1730, 828), (1420, 757))

    text_box(draw, (565, 342, 1335, 405), "Funções comuns a todos os perfis: autenticar-se, encerrar sessão, alterar senha e consultar estatísticas dentro do próprio escopo.",
             fill="F2F2F2", outline=GRAY, radius=12, padding=12, font_size=17)
    note_font = pfont(18, False)
    draw.text((550, 1100), "As permissões são verificadas novamente na API; a interface não é a única barreira de acesso.",
              font=note_font, fill=GRAY)
    return save_image(image, "figura_1_casos_uso.png")


def make_relationship_diagram() -> Path:
    image = Image.new("RGB", (1900, 1250), "WHITE")
    draw = ImageDraw.Draw(image)
    heading = pfont(31, True)
    draw.text((590, 45), "Modelo lógico de relacionamentos do SAC", font=heading, fill=BLUE)

    text_box(draw, (110, 200, 650, 540),
             "ALUNOS\nPK matrícula\nnome, curso, disciplina, turma\ncota_detalhada, descrição\nnota_final, taxa_faltas\nmaterias_json e campos legados",
             fill="FFFFFF", title=True, align="left", font_size=23)
    text_box(draw, (1230, 185, 1770, 520),
             "USUÁRIOS\nPK id | UQ username\nrole, matrícula, nome\ncurso, disciplina, turma\nsession_version\nsenha em hash bcrypt",
             fill=LIGHT_BLUE, title=True, align="left", font_size=23)
    text_box(draw, (610, 760, 1215, 1080),
             "HISTÓRICO_ACADÊMICO\nPK matrícula + período\nsnapshot_json, salvo_em, salvo_por\nRetém a fotografia do período\nmesmo se o aluno atual for excluído",
             fill="FFFFFF", title=True, align="left", font_size=24)
    text_box(draw, (1300, 760, 1770, 1035),
             "SESSÕES\nPK sid\nsess (JSON)\nexpires\nVínculo lógico: id do usuário\nserializado na sessão",
             fill="FFFFFF", title=True, align="left", font_size=23)
    text_box(draw, (100, 790, 475, 1015),
             "MOVIMENTAÇÕES\nPK id\ntipo, setor, descrição\nusuário, timestamp\nTabela prevista para\nauditoria futura",
             fill=LIGHT_GRAY, title=True, align="left", font_size=22)

    arrow(draw, (650, 360), (1230, 355), color=BLUE, width=3)
    arrow(draw, (540, 540), (730, 760), color=BLUE, width=3)
    arrow(draw, (1370, 520), (1490, 760), color=BLUE, width=3)

    rel_font = pfont(19, True)
    small = pfont(18, False)
    draw.text((748, 265), "0..1 conta Aluno por matrícula", font=rel_font, fill=BLUE)
    draw.text((720, 295), "relação lógica; índice único normalizado", font=small, fill=GRAY)
    draw.text((800, 650), "0..6 snapshots", font=rel_font, fill=BLUE)
    draw.text((1505, 630), "1..N", font=rel_font, fill=BLUE)
    draw.text((105, 1040), "Tabela prevista para auditoria futura; ainda sem fluxo de escrita/leitura integrado.", font=small, fill=GRAY)
    note = "Não há chaves estrangeiras físicas nesses vínculos: o desenho documenta relações lógicas e a retenção histórica deliberada."
    draw.text((205, 1160), note, font=pfont(20, False), fill=GRAY)
    return save_image(image, "figura_2_relacionamentos.png")


def make_architecture_diagram() -> Path:
    image = Image.new("RGB", (1900, 1120), "WHITE")
    draw = ImageDraw.Draw(image)
    heading = pfont(31, True)
    draw.text((640, 42), "Arquitetura em camadas", font=heading, fill=BLUE)
    text_box(draw, (75, 295, 430, 605),
             "NAVEGADOR\nUsuário\nAluno, Professor,\nDiretor(a) ou Admin",
             fill="FFFFFF", title=True, font_size=26)
    text_box(draw, (560, 210, 1050, 670),
             "CLIENTE WEB\nReact 19 + TypeScript + Vite\nPáginas, componentes,\nformulários, gráficos SVG\ne exportações CSV/PDF",
             fill=LIGHT_BLUE, title=True, align="left", font_size=25)
    text_box(draw, (1175, 210, 1670, 670),
             "SERVIDOR\nNode.js + Express\nRotas JSON, autenticação,\nRBAC, validações, CSV,\nsnapshots e backup",
             fill="FFFFFF", title=True, align="left", font_size=25)
    text_box(draw, (725, 820, 1450, 1040),
             "PERSISTÊNCIA\nSQLite: alunos, usuários, sessões, histórico e movimentações\nTransações imediatas para operações de cadastro e ciclo de conta",
             fill=LIGHT_BLUE, title=True, align="left", font_size=25)
    arrow(draw, (430, 450), (560, 450), color=BLUE, width=4)
    arrow(draw, (1050, 450), (1175, 450), color=BLUE, width=4)
    arrow(draw, (1420, 670), (1120, 820), color=BLUE, width=4)
    label_font = pfont(19, True)
    draw.text((445, 405), "HTTPS/HTTP", font=label_font, fill=GRAY)
    draw.text((1062, 405), "/api", font=label_font, fill=GRAY)
    draw.text((1200, 740), "consultas e transações", font=label_font, fill=GRAY)
    draw.text((570, 1055), "O Express também entrega os artefatos estáticos compilados do React na execução simplificada.",
              font=pfont(19, False), fill=GRAY)
    return save_image(image, "figura_3_arquitetura.png")


def set_run_font(run, size=12, bold=None, italic=None, color=INK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag, attributes in [
        ("w:tblW", {qn("w:w"): str(total), qn("w:type"): "dxa"}),
        ("w:tblInd", {qn("w:w"): "0", qn("w:type"): "dxa"}),
    ]:
        existing = tbl_pr.find(qn(tag))
        if existing is None:
            existing = OxmlElement(tag)
            tbl_pr.append(existing)
        for key, value in attributes.items():
            existing.set(key, value)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for grid_col in list(grid):
        grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 10, color=INK)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc_field(paragraph):
    """Insert a Word-native TOC field; Word updates it automatically on open."""
    run = paragraph.add_run()
    set_run_font(run, 12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Atualize o campo de sumário no Microsoft Word, se necessário."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, fallback, end])


def add_figure_list_field(paragraph):
    """Insert a Word-native list-of-figures field based on SEQ Figura captions."""
    run = paragraph.add_run()
    set_run_font(run, 12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\h \\z \\c "Figura" '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Atualize a lista de figuras no Microsoft Word, se necessário."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, fallback, end])


def add_sequence_field(run, sequence_name, fallback_number):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {sequence_name} \\* ARABIC "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = str(fallback_number)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number_section(section):
    section.header.is_linked_to_previous = False
    section.header_distance = Cm(2)
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_page_field(paragraph)
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is None:
        existing = OxmlElement("w:pgNumType")
        sect_pr.append(existing)
    existing.set(qn("w:start"), "1")


def set_image_alt(inline, title, description):
    doc_prs = inline._inline.xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[0].set("title", title)
        doc_prs[0].set("descr", description)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(2)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)

    h1 = styles["Heading 1"]
    h1.font.name = FONT
    h1._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(24)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.first_line_indent = Cm(0)

    h2 = styles["Heading 2"]
    h2.font.name = FONT
    h2._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.space_before = Pt(24)
    h2.paragraph_format.space_after = Pt(12)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.first_line_indent = Cm(0)

    h3 = styles["Heading 3"]
    h3.font.name = FONT
    h3._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h3._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h3.font.size = Pt(12)
    h3.font.bold = False
    h3.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.line_spacing = 1.5
    h3.paragraph_format.space_before = Pt(18)
    h3.paragraph_format.space_after = Pt(8)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.first_line_indent = Cm(0)

    if "Figure Caption" not in [style.name for style in styles]:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(10)
    caption.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(0)
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Source Note" not in [style.name for style in styles]:
        source = styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = styles["Source Note"]
    source.font.name = FONT
    source._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    source._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    source.font.size = Pt(10)
    source.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
    source.paragraph_format.line_spacing = 1.0
    source.paragraph_format.space_before = Pt(0)
    source.paragraph_format.space_after = Pt(18)
    source.paragraph_format.first_line_indent = Cm(0)
    source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_direct_paragraph(doc, text="", *, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=False,
                         italic=False, before=0, after=0, line=1.5, first_indent=0, color=INK):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.first_line_indent = Cm(first_indent)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    set_run_font(run, 12)
    return paragraph


def add_h1(doc, title, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title.upper())
    return p


def add_h2(doc, title):
    p = doc.add_paragraph(style="Heading 2")
    p.add_run(title)
    return p


def add_h3(doc, title):
    p = doc.add_paragraph(style="Heading 3")
    p.add_run(title)
    return p


def add_figure(doc, image_path: Path, figure_number: int, title: str, description: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Cm(15.5))
    set_image_alt(inline, f"Figura {figure_number} – {title}", description)
    c = doc.add_paragraph(style="Figure Caption")
    prefix = c.add_run("Figura ")
    set_run_font(prefix, 10)
    number_run = c.add_run()
    set_run_font(number_run, 10)
    add_sequence_field(number_run, "Figura", figure_number)
    suffix = c.add_run(f" – {title}")
    set_run_font(suffix, 10)
    s = doc.add_paragraph(style="Source Note")
    s.add_run("Fonte: elaboração própria (2026).")


def add_table(doc, headers, rows, widths, source=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "D9E5F3")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(heading)
        set_run_font(run, 10, bold=True)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.first_line_indent = Cm(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, 10)
    set_table_geometry(table, widths)
    if source:
        source_p = doc.add_paragraph(style="Source Note")
        source_p.add_run(source)
    else:
        spacer = doc.add_paragraph(style="Source Note")
        spacer.add_run("")
    return table


def add_front_matter_title(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(24)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    set_run_font(run, 14, bold=True)
    return paragraph


def dotted_line(title, page):
    left = title
    dots = "." * max(3, 92 - len(left) - len(str(page)))
    return f"{left} {dots} {page}"


def add_cover(doc):
    for line in [
        "MINISTÉRIO DA EDUCAÇÃO",
        "SECRETARIA DE EDUCAÇÃO PROFISSIONAL E TECNOLÓGICA",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA",
        "[CAMPUS / DIRETORIA]",
        "[CURSO]",
    ]:
        add_direct_paragraph(doc, line, size=14, bold=True, after=0)
    add_direct_paragraph(doc, "", before=72)
    add_direct_paragraph(doc, "[NOME DO(A) AUTOR(A)]", size=16, bold=True)
    add_direct_paragraph(doc, "", before=92)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO DE DESENVOLVIMENTO E EVOLUÇÃO", size=18, bold=True)
    add_direct_paragraph(doc, "DO SISTEMA DE ALUNOS COTISTAS (SAC)", size=18, bold=True)
    add_direct_paragraph(doc, "", before=70)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO/CIENTÍFICO", size=14, bold=True)
    add_direct_paragraph(doc, "", before=78)
    add_direct_paragraph(doc, "[CIDADE - UF]", size=14, bold=True)
    add_direct_paragraph(doc, "2026", size=14, bold=True)


def add_title_page(doc):
    doc.add_page_break()
    add_direct_paragraph(doc, "[NOME DO(A) AUTOR(A)]", size=14, bold=True, before=0)
    add_direct_paragraph(doc, "", before=86)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO DE DESENVOLVIMENTO E EVOLUÇÃO", size=16, bold=True)
    add_direct_paragraph(doc, "DO SISTEMA DE ALUNOS COTISTAS (SAC)", size=16, bold=True)
    add_direct_paragraph(doc, "", before=62)
    natureza = (
        "Relatório técnico/científico elaborado como documento de apoio ao Trabalho de Conclusão de Curso, "
        "apresentado ao [CURSO] do Instituto Federal de Educação, Ciência e Tecnologia de Roraima, "
        "como requisito parcial para obtenção do título de [GRAU].\n\n"
        "Orientador(a): [NOME DO(A) ORIENTADOR(A)]"
    )
    add_direct_paragraph(doc, natureza, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    add_direct_paragraph(doc, "", before=115)
    add_direct_paragraph(doc, "[CIDADE - UF]", size=14, bold=False)
    add_direct_paragraph(doc, "2026", size=14, bold=False)


def add_pre_textual(doc, toc, figures):
    doc.add_page_break()
    add_front_matter_title(doc, "RESUMO")
    resumo = (
        "Este relatório técnico/científico apresenta a engenharia, as funcionalidades e as alterações "
        "consolidadas no Sistema de Alunos Cotistas (SAC), uma aplicação web destinada ao acompanhamento "
        "acadêmico de estudantes cotistas. O trabalho descreve a migração da interface para React com "
        "TypeScript, a organização do servidor em Node.js e Express, a persistência em SQLite e a "
        "implementação de controle de acesso por papéis. São detalhados o tratamento de notas e faltas, "
        "a situação acadêmica de risco, o cadastro manual e por arquivo CSV, as estatísticas dinâmicas, "
        "o armazenamento de históricos por período, as exportações e a autenticação. A abordagem combina "
        "inspeção técnica do código, análise das regras de negócio e execução de verificações automatizadas. "
        "Os resultados apontam uma solução organizada em camadas, com regras de escopo aplicadas no servidor, "
        "contas estudantis vinculadas ao cadastro, sessões persistentes e mecanismos de integridade para "
        "evitar conflitos de matrícula. O documento também registra limitações observadas e recomendações "
        "para evolução segura do sistema em contexto institucional."
    )
    p = add_direct_paragraph(doc, resumo, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    p.paragraph_format.space_after = Pt(12)
    add_direct_paragraph(doc, "Palavras-chave: sistema acadêmico; cotas; autenticação; React; estatísticas.",
                         align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)

    doc.add_page_break()
    add_front_matter_title(doc, "ABSTRACT")
    abstract = (
        "This technical/scientific report presents the engineering, features and consolidated changes of "
        "the Sistema de Alunos Cotistas (SAC), a web application for the academic monitoring of quota "
        "students. It describes the migration of the interface to React with TypeScript, the Node.js and "
        "Express server organization, SQLite persistence and role-based access control. The report covers "
        "grade and attendance processing, academic risk status, manual and CSV registration, dynamic "
        "statistics, period snapshots, exports and authentication. The approach combines code inspection, "
        "business-rule analysis and automated verification. The resulting solution is organized in layers, "
        "enforces scopes on the server, links student accounts to enrollment records, persists sessions and "
        "prevents enrollment identifier conflicts. Limitations and recommendations for a safer institutional "
        "evolution are also documented."
    )
    add_direct_paragraph(doc, abstract, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    add_direct_paragraph(doc, "Keywords: academic system; quotas; authentication; React; statistics.",
                         align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0, before=12)

    doc.add_page_break()
    add_front_matter_title(doc, "LISTA DE FIGURAS")
    figures_paragraph = doc.add_paragraph()
    figures_paragraph.paragraph_format.first_line_indent = Cm(0)
    figures_paragraph.paragraph_format.line_spacing = 1.5
    add_figure_list_field(figures_paragraph)

    doc.add_page_break()
    add_front_matter_title(doc, "LISTA DE SIGLAS")
    siglas = [
        ("API", "Application Programming Interface"),
        ("CSV", "Comma-Separated Values"),
        ("IFRR", "Instituto Federal de Educação, Ciência e Tecnologia de Roraima"),
        ("RBAC", "Role-Based Access Control"),
        ("SAC", "Sistema de Alunos Cotistas"),
        ("SPA", "Single Page Application"),
        ("TCC", "Trabalho de Conclusão de Curso"),
    ]
    for sigla, meaning in siglas:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(f"{sigla} – ")
        set_run_font(run, 12, bold=True)
        run = p.add_run(meaning)
        set_run_font(run, 12)

    doc.add_page_break()
    add_front_matter_title(doc, "SUMÁRIO")
    toc_paragraph = doc.add_paragraph()
    toc_paragraph.paragraph_format.first_line_indent = Cm(0)
    toc_paragraph.paragraph_format.line_spacing = 1.5
    add_toc_field(toc_paragraph)


def add_intro(doc):
    add_h1(doc, "1 INTRODUÇÃO", page_break=False)
    add_body(doc, "O acompanhamento acadêmico de estudantes cotistas exige a organização de dados de matrícula, curso, turma, disciplina, notas, faltas e informações de apoio. Quando esses dados são mantidos em fluxos dispersos, a identificação de situações de atenção e a tomada de decisão ficam mais lentas. O Sistema de Alunos Cotistas (SAC) foi consolidado para concentrar esse acompanhamento em uma aplicação web com acesso diferenciado para estudante, professor, direção e administração.")
    add_body(doc, "O presente documento caracteriza o produto como relatório técnico/científico de apoio ao TCC. Essa modalidade é prevista pelo IFRR, e a estrutura adotada combina introdução, desenvolvimento, considerações finais, referências, apêndices e elementos pré-textuais. As regras de apresentação empregadas neste arquivo seguem o Manual de Normas para Elaboração de Trabalhos Acadêmicos do IFRR e devem ser conferidas pelo autor e pelo orientador em relação ao PPC do curso antes da submissão final.")
    add_h2(doc, "1.1 Problema e objetivo")
    add_body(doc, "O problema abordado é a necessidade de registrar e consultar informações acadêmicas de estudantes cotistas de forma estruturada, preservando a confidencialidade, o escopo de cada perfil de acesso e a evolução entre períodos. O objetivo geral foi organizar um sistema web que apoie o cadastro, a visualização acadêmica, a identificação de risco, as estatísticas e a administração de contas da equipe, com código de complexidade moderada e separação clara de responsabilidades.")
    add_h2(doc, "1.2 Escopo do relatório")
    add_body(doc, "O relatório descreve a versão entregue do SAC, com foco na arquitetura, nas regras implementadas, no modelo lógico de dados, na autenticação, no controle de acesso, na importação de arquivos, nos históricos, nas estatísticas e nas verificações realizadas. Não é uma certificação de conformidade jurídica, de segurança ou de proteção de dados; recomendações complementares são apresentadas como oportunidades de melhoria.")


def add_method(doc):
    add_h1(doc, "2 MÉTODO DE DESENVOLVIMENTO E DELIMITAÇÃO")
    add_body(doc, "A evolução do sistema foi conduzida por incrementos orientados por funcionalidades: primeiro a organização da interface e do fluxo de acesso; depois o cadastro e a consulta de alunos; em seguida, a análise acadêmica, os históricos e os relatórios; por fim, a revisão de integridade das contas e das permissões. Essa estratégia permitiu validar cada regra perto do ponto em que ela é aplicada, evitando concentrar toda a lógica no cliente web.")
    add_h2(doc, "2.1 Procedimentos técnicos")
    add_body(doc, "Foram realizadas inspeções dos módulos do projeto, leitura das rotas e dos contratos de dados, execução dos comandos de teste, compilação e análise estática do cliente. A análise também verificou o comportamento do cadastro manual e da importação CSV, a vinculação de contas de alunos, a autenticação por identificadores permitidos e a aplicação do escopo de visibilidade no servidor.")
    add_h2(doc, "2.2 Critérios de engenharia")
    add_body(doc, "Os critérios adotados foram separação entre interface, rotas, regras de domínio e persistência; validação de entrada; transações para operações acopladas; autenticação com senha em hash; sessão persistente; autorização por papel; preservação de snapshots acadêmicos; e retorno padronizado de erros da API. O código foi mantido com módulos pequenos e comentários de propósito nos arquivos principais, reduzindo a dependência de abstrações excessivas.")
    add_h2(doc, "2.3 Referencial tecnológico")
    add_body(doc, "A interface utiliza React e TypeScript, estrutura adequada para componentes, estados e contratos de dados. O servidor utiliza Express para expor uma API JSON e servir a versão compilada da aplicação. SQLite foi escolhido como persistência embarcada, e as operações que exigem consistência, como a criação conjunta de aluno e conta, são agrupadas em transações. A documentação oficial dessas tecnologias e as recomendações da OWASP foram usadas como referência técnica complementar.")


def add_overview(doc):
    add_h1(doc, "3 VISÃO GERAL DO SISTEMA")
    add_body(doc, "O SAC é uma aplicação web de página única. O navegador carrega a interface React, que solicita dados à API Express. A API aplica autenticação, validação e autorização antes de consultar ou alterar o banco SQLite. Essa organização permite que a navegação da interface seja simples, sem transferir para o navegador decisões de segurança ou filtros de escopo.")
    add_h2(doc, "3.1 Perfis e escopos")
    add_table(doc,
              ["Perfil", "Escopo de visualização", "Ações principais"],
              [
                  ("Aluno", "Somente a própria matrícula", "Consulta seu desempenho, estatísticas pessoais e altera a senha."),
                  ("Professor", "Curso + disciplina + turma atribuídos", "Consulta e gerencia estudantes de sua turma; utiliza estatísticas do respectivo escopo."),
                  ("Diretor(a)", "Curso atribuído", "Acompanha e gerencia estudantes do curso sob sua responsabilidade."),
                  ("Administrador(a)", "Todos os registros", "Administra equipe, filtros institucionais, histórico por período e backup."),
              ],
              [1600, 2700, 4740],
              "Fonte: elaboração própria, com base na implementação do SAC (2026).")
    add_body(doc, "A regra de escopo é aplicada no servidor por funções de autorização. Desse modo, ocultar uma opção na interface não é considerado controle suficiente: as rotas também verificam o perfil autenticado e os campos de curso, disciplina, turma ou matrícula antes de entregar ou aceitar dados.")
    add_h2(doc, "3.2 Situação acadêmica")
    add_body(doc, "Cada aluno pode possuir até vinte matérias, com nota, faltas, faltas justificadas, total de aulas e situação por componente. O sistema calcula indicadores agregados de nota e frequência e mantém a classificação Regular, Alerta ou Em Risco. A lógica preexistente de nota e falta foi preservada e adaptada para a estrutura de matérias, sem diluir o resumo acadêmico anterior.")
    add_h2(doc, "3.3 Navegação e experiência de uso")
    add_body(doc, "A aplicação apresenta painel principal, estatísticas, cadastro, matérias, períodos, configurações e permissões administrativas. A identidade visual adota o nome Sistema de Alunos Cotistas (SAC), com referência institucional do IFRR, e oferece tema claro ou escuro. Quando um aluno possui descrição registrada, seu nome pode ser focalizado e acionado para mostrar uma prévia e abrir uma janela acessível com o texto completo.")


def add_engineering(doc):
    add_h1(doc, "4 ENGENHARIA E ALTERAÇÕES IMPLEMENTADAS")
    add_h2(doc, "4.1 Migração e organização da interface")
    add_body(doc, "A interface foi organizada em React 19, TypeScript e Vite. As páginas concentram os fluxos de tela, os componentes reutilizáveis representam formulários, tabelas e gráficos, e os utilitários concentram chamadas à API, contratos de tipos, conversão de dados e geração de PDF. A navegação usa rotas por hash, o que simplifica o uso da SPA servida pelo próprio Express sem introduzir dependência adicional de roteamento.")
    add_h2(doc, "4.2 Estrutura do servidor e da API")
    add_body(doc, "O servidor foi dividido em entrada da aplicação, rotas de autenticação, rotas de domínio, serviços de dados, permissões, utilitários, respostas HTTP e armazenamento de sessões. As rotas agrupam endpoints de alunos, períodos, estatísticas, usuários de equipe, backup e importação. A camada de dados encapsula consultas SQL e operações transacionais, deixando a camada HTTP responsável por validar a requisição e retornar respostas consistentes.")
    add_h2(doc, "4.3 Cadastro, matérias e importação de arquivos")
    add_body(doc, "O cadastro manual valida matrícula, dados acadêmicos e matérias. A importação CSV recebe arquivo em memória, limita o tamanho a 1 MB e mil linhas, exige cabeçalho conhecido, valida as colunas e informa linhas importadas, ignoradas ou com erro. Também existe modelo CSV para download. A matrícula é normalizada com remoção de espaços externos e padronização de caixa, e um índice único case-insensitive impede que variações como letras maiúsculas e minúsculas representem alunos diferentes.")
    add_h2(doc, "4.4 Estatísticas, gráficos e históricos")
    add_body(doc, "A página de estatísticas calcula totais, distribuição de notas, distribuição de frequência, situação de risco e médias por matéria dentro do escopo do usuário. O usuário pode alternar entre cinco visualizações SVG: pizza, colunas, linhas, barras e área. Administradores recebem filtros adicionais por curso, turma, categoria acadêmica e modalidade de cota; os demais perfis não recebem nem enviam esses filtros.")
    add_body(doc, "Para permitir comparações ao longo do ano, o administrador pode salvar um snapshot por período. O sistema trabalha com quatro bimestres e dois semestres. Depois de salvo, um período usa somente o conjunto de alunos capturado no snapshot, evitando que alunos cadastrados depois alterem indevidamente dados históricos. A atualização exige confirmação explícita; estudantes removidos do cadastro atual permanecem recuperáveis no retrato histórico.")
    add_h2(doc, "4.5 Exportação, impressão e inicialização")
    add_body(doc, "As estatísticas podem ser exportadas em CSV, geradas em PDF formatado e impressas. O PDF é produzido no cliente com carregamento dinâmico da biblioteca de geração, evitando aumentar desnecessariamente o carregamento inicial da aplicação. Para a operação local, o comando quickstart compila o cliente, inicia o servidor e abre a URL correspondente no navegador. A criação inicial do administrador é feita por script separado, sem expor credenciais em código-fonte.")


def add_diagrams(doc, figures):
    add_h1(doc, "5 MODELAGEM E DIAGRAMAS")
    add_h2(doc, "5.1 Casos de uso")
    add_body(doc, "A Figura 1 apresenta os principais atores e casos de uso. O núcleo comum é a autenticação e a alteração de senha. A diferença entre perfis não está apenas na tela disponível, mas no recorte de dados que cada ator pode consultar ou modificar. A administração reúne atividades de manutenção institucional, enquanto estudantes acessam exclusivamente dados relacionados à sua própria matrícula.")
    add_figure(doc, figures["use_case"], 1, "Diagrama de casos de uso do Sistema de Alunos Cotistas (SAC)",
               "Diagrama UML simplificado com os atores Aluno, Professor, Diretor(a) e Administrador(a), conectados às funções principais do SAC.")
    add_h2(doc, "5.2 Relacionamentos de dados")
    add_body(doc, "A Figura 2 representa o modelo lógico. O cadastro do aluno é a referência para a conta de papel Aluno, para os snapshots acadêmicos e para o escopo das consultas. Algumas relações são propositalmente lógicas, e não chaves estrangeiras físicas: o histórico deve sobreviver à exclusão do aluno atual e a sessão guarda o identificador do usuário em conteúdo serializado. Essa escolha é documentada para que futuras evoluções preservem o comportamento esperado.")
    add_figure(doc, figures["relationship"], 2, "Diagrama lógico de relacionamentos de dados",
               "Modelo lógico das tabelas alunos, usuários, histórico acadêmico, sessões e movimentações, com relações e notas de implementação.")
    add_h2(doc, "5.3 Arquitetura")
    add_body(doc, "A Figura 3 mostra o particionamento em camadas. O navegador opera a interface, a SPA organiza a experiência e solicita recursos da API, o servidor concentra regras de domínio e segurança, e o banco persiste o estado. A configuração favorece uma implantação local simples, mantendo uma separação suficiente para testes, manutenção e futura substituição de componentes.")
    add_figure(doc, figures["architecture"], 3, "Arquitetura em camadas do SAC",
               "Fluxo entre navegador, cliente React, servidor Express e banco SQLite, incluindo as responsabilidades principais de cada camada.")


def add_security(doc):
    add_h1(doc, "6 SEGURANÇA, AUTENTICAÇÃO E CONTROLE DE ACESSO")
    add_h2(doc, "6.1 Autenticação e senhas")
    add_body(doc, "As senhas são persistidas como hash bcrypt. A aplicação valida o tamanho da nova senha, exige confirmação e incrementa uma versão de sessão quando a senha é alterada, invalidando sessões antigas. As contas de aluno são criadas automaticamente no cadastro manual e na importação, na mesma transação do aluno. A regra de senha inicial legada foi mantida como procedimento de primeira entrada, porém seu valor não é armazenado em texto puro e não é exposto neste relatório por se tratar de informação operacional sensível.")
    add_body(doc, "O login aceita identificadores coerentes com o perfil: usuário, matrícula normalizada ou nome completo quando ele é único para contas permitidas. Consultas por nome ambíguo são recusadas em vez de selecionar um usuário de maneira arbitrária. A aplicação corrige perfis de conta que estejam divergentes do cadastro atual sem substituir senha personalizada; somente credenciais comprovadamente ainda pertencentes à regra inicial podem ser redefinidas durante uma reconciliação controlada.")
    add_h2(doc, "6.2 Sessões e proteção de requisições")
    add_body(doc, "As sessões são armazenadas em SQLite, têm expiração e utilizam cookie HTTP-only, SameSite=Lax e atributo Secure em produção. O servidor renova a sessão após login e bloqueia mudanças de estado oriundas de origem cruzada não esperada. Tentativas inválidas de acesso são limitadas por combinação de endereço IP e identificador normalizado, reduzindo o efeito de erros de um usuário sobre os demais usuários da mesma rede.")
    add_h2(doc, "6.3 Ciclo de vida de contas estudantis")
    add_body(doc, "Foi implementado um mecanismo de limpeza para contas antigas de papel Aluno que não possuem matrícula ativa correspondente. A limpeza ocorre na inicialização e antes da criação de um novo aluno, remove as sessões da conta órfã e preserva contas da equipe e snapshots acadêmicos. Um gatilho também revoga a conta do aluno quando o cadastro é excluído. A tela genérica de permissões deixou de criar ou converter contas de alunos, o que impede a transferência indevida de senha ou identidade entre estudantes.")
    add_table(doc,
              ["Controle", "Tratamento implementado", "Finalidade"],
              [
                  ("Matrícula", "Normalização e índice único sem diferenciar maiúsculas/minúsculas.", "Evitar colisão e ambiguidade de login."),
                  ("Conta de aluno", "Criação transacional, backfill e reconciliação de perfil.", "Manter cadastro e identidade de acesso consistentes."),
                  ("Senha", "bcrypt, troca autenticada e versionamento de sessão.", "Reduzir exposição e revogar sessões antigas."),
                  ("Autorização", "RBAC e verificação de escopo na API.", "Impedir acesso fora da matrícula, turma ou curso."),
                  ("Conta órfã", "Limpeza seletiva e remoção de sessões associadas.", "Evitar reutilização indevida de identidade antiga."),
              ],
              [1700, 4000, 3340],
              "Fonte: elaboração própria, com base na implementação do SAC e em recomendações de segurança de autenticação (OWASP, 2026).")


def add_validation(doc):
    add_h1(doc, "7 VALIDAÇÃO E QUALIDADE")
    add_body(doc, "A validação foi conduzida em proporção ao risco das alterações. O objetivo não foi somente compilar a aplicação, mas testar os pontos de integridade em que uma inconsistência poderia impedir acesso de aluno, reaproveitar uma conta antiga ou misturar dados de períodos diferentes. As verificações foram realizadas sem expor credenciais, dados pessoais ou cópias do banco institucional no relatório.")
    add_h2(doc, "7.1 Verificações executadas")
    add_table(doc,
              ["Verificação", "Evidência", "Resultado"],
              [
                  ("Testes de ciclo de conta", "node --test test/student-account-cleanup.test.js", "Três cenários aprovados: limpeza, sessão e reutilização controlada."),
                  ("Compilação do cliente", "npm run build", "TypeScript e Vite concluídos sem erro."),
                  ("Análise estática", "npm --prefix teste-react run lint", "Verificação do cliente concluída sem erro."),
                  ("Sintaxe do servidor", "node --check em módulos principais", "Módulos verificados sem erro de sintaxe."),
                  ("Dependências", "npm audit --omit=dev", "Auditoria sem vulnerabilidades reportadas no momento da validação."),
              ],
              [1900, 3750, 3390],
              "Fonte: elaboração própria (2026).")
    add_h2(doc, "7.2 Cenários de integração relevantes")
    add_body(doc, "Os cenários mais sensíveis incluíram criação manual de aluno com conta associada, importação CSV com conta associada, login por matrícula e por nome único, preservação de senha personalizada, recusa de colisão entre matrícula e username, restrição de tentativas por identificador, fechamento de período sem inclusão retroativa de aluno novo e recuperação de aluno removido em snapshot salvo. Esses testes confirmam o comportamento esperado para os fluxos que motivaram as correções de autenticação e de histórico.")
    add_h2(doc, "7.3 Critério de aceitação")
    add_body(doc, "A aceitação técnica da versão considera que o cliente compila, o lint não aponta problemas, os testes de ciclo de conta passam e os fluxos de acesso e histórico mantêm a consistência descrita. A aceitação institucional deve ainda considerar homologação por usuários responsáveis, validação das regras acadêmicas e configuração de infraestrutura de produção, especialmente TLS, segredo de sessão e rotina de backup.")


def add_limitations(doc):
    add_h1(doc, "8 LIMITAÇÕES E OPORTUNIDADES DE MELHORIA")
    add_body(doc, "A versão atual atende ao escopo funcional consolidado, mas uma avaliação técnica identificou evoluções recomendadas antes de uma implantação institucional ampliada. Essas recomendações não invalidam os controles existentes; elas tornam o produto mais auditável, resiliente e alinhado a operações de longo prazo.")
    add_table(doc,
              ["Prioridade", "Melhoria sugerida", "Justificativa"],
              [
                  ("Alta", "Completar o README com arquitetura, variáveis, perfis, importação e manutenção.", "O repositório possui QUICKSTART, mas a documentação principal ainda é insuficiente para repasse de conhecimento."),
                  ("Alta", "Configurar HTTPS, SESSION_SECRET forte e backup externo/versionado em produção.", "A segurança do transporte e a recuperação de dados dependem da infraestrutura de implantação."),
                  ("Alta", "Ampliar testes de API, RBAC, importação, login e interface.", "A suíte atual cobre de modo aprofundado apenas o ciclo de contas órfãs."),
                  ("Média", "Integrar movimentações como trilha de auditoria consultável.", "A tabela existe, porém ainda não é alimentada nem lida pelo fluxo da aplicação."),
                  ("Média", "Implementar rotação configurável de backups e teste de restauração.", "O backup atual é uma cópia única e não comprova a recuperação completa."),
                  ("Média", "Validar formalmente parâmetros de risco com a área pedagógica.", "Regras de nota e frequência exigem governança acadêmica e documentação institucional."),
                  ("Média", "Substituir o limitador em memória por armazenamento compartilhado em escala.", "Em múltiplos processos ou reinicializações, o limite não é distribuído nem persistente."),
                  ("Evolutiva", "Adicionar MFA, recuperação de senha assistida e registro de eventos de segurança.", "Fortalece contas sensíveis e melhora rastreabilidade em cenários institucionais."),
              ],
              [1100, 4000, 3940],
              "Fonte: elaboração própria (2026).")
    add_body(doc, "Também se recomenda documentar a decisão de manter certos vínculos como relações lógicas, pois o histórico precisa sobreviver à exclusão de um aluno ativo. Caso o banco evolua para um SGBD externo, essa regra deverá ser preservada por política de retenção e não por exclusão em cascata indiscriminada.")


def add_conclusion(doc):
    add_h1(doc, "9 CONSIDERAÇÕES FINAIS")
    add_body(doc, "O Sistema de Alunos Cotistas (SAC) foi estruturado como uma aplicação web com fronteiras claras entre interface, API, regras de acesso e persistência. A migração para React e TypeScript tornou a camada visual mais modular, enquanto Express e SQLite concentram a lógica de domínio e os dados. A solução cobre cadastro, matérias, notas, faltas, situação de risco, estatísticas, snapshots acadêmicos, importação e exportação de informações.")
    add_body(doc, "O principal ganho de engenharia está na combinação entre escopo de acesso verificado no servidor, criação transacional de conta estudantil, normalização de matrícula, persistência de sessões, reconciliação controlada de perfis e preservação histórica por período. Essas decisões tratam tanto da experiência de uso quanto da integridade dos registros, especialmente em situações de cadastro novo, reutilização de matrícula e comparação de desempenho ao longo do ano.")
    add_body(doc, "Para uso como TCC, o autor deve preencher os campos de identificação da capa e folha de rosto, confirmar os elementos exigidos pelo PPC e pelo orientador, atualizar o sumário após qualquer alteração de paginação e revisar as referências de acordo com as versões vigentes das normas ABNT disponibilizadas pelo IFRR. O relatório fornece uma base técnica organizada, acompanhada de diagramas e recomendações para a continuidade do sistema.")


def add_references(doc):
    add_h1(doc, "REFERÊNCIAS")
    references = [
        "EXPRESSJS. Production best practices: security. [S. l.], 2026. Disponível em: <https://expressjs.com/en/advanced/best-practice-security/>. Acesso em: 17 ago. 2026.",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA. Manual de normas para elaboração de trabalhos acadêmicos do IFRR. Boa Vista: IFRR, [s. d.]. Disponível em: <https://boavista.ifrr.edu.br/aluno/manual-detrabalhos-academicos-do-ifrr/manual-de-trabalhos-academicos-do-ifrr/at_download/file>. Acesso em: 17 ago. 2026.",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA. Resolução CONSUP/IFRR nº 730, de 30 de março de 2023. Estabelece normas e diretrizes para a elaboração do Trabalho de Conclusão de Curso (TCC) dos cursos de graduação, no âmbito do Instituto Federal de Roraima. Boa Vista: IFRR, 2023. Disponível em: <https://www.ifrr.edu.br/documents/1566/Resolução_n._730_-2023.pdf>. Acesso em: 17 ago. 2026.",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA. Catálogo de normas da ABNT. Boa Vista: IFRR, [s. d.]. Disponível em: <https://boavista.ifrr.edu.br/biblioteca/catalogo-de-normas-da-abnt/>. Acesso em: 17 ago. 2026.",
        "OPEN WEB APPLICATION SECURITY PROJECT. Authentication cheat sheet. [S. l.], 2026. Disponível em: <https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html>. Acesso em: 17 ago. 2026.",
        "REACT. React documentation. [S. l.], 2026. Disponível em: <https://react.dev/>. Acesso em: 17 ago. 2026.",
        "SISTEMA DE ALUNOS COTISTAS. Código-fonte da aplicação: versão analisada. [S. l.]: repositório do projeto, 2026.",
        "SQLITE. Transaction. [S. l.], 2026. Disponível em: <https://www.sqlite.org/lang_transaction.html>. Acesso em: 17 ago. 2026.",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(reference)
        set_run_font(run, 12)


def add_appendices(doc):
    add_h1(doc, "APÊNDICE A - MATRIZ DE FUNCIONALIDADES")
    add_body(doc, "A matriz a seguir relaciona funcionalidades da versão analisada a componentes técnicos que oferecem evidência de sua implementação. Os caminhos são apresentados para facilitar a continuidade da manutenção, não como substitutos da leitura integral do código.")
    add_table(doc,
              ["Funcionalidade", "Evidência técnica", "Componente principal"],
              [
                  ("Sessão e login", "Rotas de login, logout e senha; store SQLite.", "src/routes/auth.js; src/lib/auth.js; src/lib/sessionStore.js"),
                  ("RBAC", "Escopo por matrícula, curso, disciplina e turma.", "src/lib/permissions.js; src/lib/helpers.js"),
                  ("Cadastro e conta", "Transação, normalização e limpeza de contas órfãs.", "src/lib/database.js; src/routes/api.js"),
                  ("Importação CSV", "Upload em memória, validações e modelo para download.", "src/routes/api.js; pages/RegistrationPage.tsx"),
                  ("Estatísticas", "Agregação, filtros, gráficos SVG e exportação.", "pages/StatisticsPage.tsx; components/StatisticsChart.tsx"),
                  ("Histórico", "Snapshot por período e evolução anual.", "src/lib/database.js; src/routes/api.js"),
                  ("Descrição do aluno", "Tooltip e modal acessível.", "components/StudentTable.tsx"),
                  ("Operação local", "Compilação, servidor e abertura no navegador.", "src/scripts/quickstart.js; QUICKSTART.md"),
              ],
              [2000, 3600, 3440],
              "Fonte: elaboração própria, com base no repositório do SAC (2026).")

    add_h1(doc, "APÊNDICE B - ROTEIRO DE OPERAÇÃO E VERIFICAÇÃO")
    add_body(doc, "Antes de usar o sistema em uma estação local, deve-se instalar as dependências do servidor e do cliente. Em seguida, o comando quickstart compila o React, inicia o Express e abre a página no navegador. Para uma base vazia, a criação do primeiro administrador exige variáveis de ambiente próprias, evitando credencial fixa no repositório.")
    add_table(doc,
              ["Etapa", "Comando ou ação", "Verificação esperada"],
              [
                  ("1", "npm install", "Dependências do servidor instaladas."),
                  ("2", "npm --prefix teste-react install", "Dependências do cliente instaladas."),
                  ("3", "npm run quickstart", "Cliente compilado, servidor iniciado e navegador aberto."),
                  ("4", "npm test", "Testes de ciclo de contas aprovados."),
                  ("5", "npm run build e npm --prefix teste-react run lint", "Cliente compilado e analisado antes de entrega."),
                  ("6", "Homologação por perfil", "Aluno, Professor, Diretor(a) e Admin verificam seus próprios escopos."),
              ],
              [800, 4100, 4140],
              "Fonte: elaboração própria, com base nos scripts do projeto (2026).")
    add_body(doc, "Em produção, deve-se fornecer SESSION_SECRET forte, operar atrás de HTTPS, proteger cópias do banco e validar periodicamente a restauração de backup. A liberação para usuários finais deve ocorrer após homologação acadêmica das regras de risco e das modalidades de cota.")


def read_mapping(path: str | None):
    if not path:
        return {}, {}
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    return raw.get("toc", {}), raw.get("figures", {})


def build_document(mapping_file: str | None = None):
    toc, figure_pages = read_mapping(mapping_file)
    figures = {
        "use_case": make_use_case_diagram(),
        "relationship": make_relationship_diagram(),
        "architecture": make_architecture_diagram(),
    }
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Relatório Técnico do Sistema de Alunos Cotistas (SAC)"
    doc.core_properties.subject = "Documento de apoio a TCC"
    doc.core_properties.author = "Sistema de Alunos Cotistas"
    doc.core_properties.comments = ""
    doc.core_properties.keywords = "SAC, IFRR, TCC, sistema acadêmico"
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)

    add_cover(doc)
    add_title_page(doc)
    add_pre_textual(doc, toc, figure_pages)

    text_section = doc.add_section(WD_SECTION.NEW_PAGE)
    text_section.page_width = Cm(21)
    text_section.page_height = Cm(29.7)
    text_section.top_margin = Cm(3)
    text_section.bottom_margin = Cm(2)
    text_section.left_margin = Cm(3)
    text_section.right_margin = Cm(2)
    add_page_number_section(text_section)

    add_intro(doc)
    add_method(doc)
    add_overview(doc)
    add_engineering(doc)
    add_diagrams(doc, figures)
    add_security(doc)
    add_validation(doc)
    add_limitations(doc)
    add_conclusion(doc)
    add_references(doc)
    add_appendices(doc)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--mapping", help="JSON with final table-of-contents and figure pages")
    args = parser.parse_args()
    build_document(args.mapping)
