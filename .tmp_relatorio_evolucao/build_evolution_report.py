from __future__ import annotations

"""Create the consolidated technical-evolution report for the SAC project.

The report is authored from the supplied legacy report and the current technical
report/source tree. It deliberately labels historical claims as historical and
does not turn roadmap items from the legacy document into implemented features.
"""

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
LEGACY_BUILDER_DIR = ROOT / ".tmp_relatorio_tcc"
sys.path.insert(0, str(LEGACY_BUILDER_DIR))
TABLE_HELPER_DIR = Path(r"C:\Users\blrsj\.codex\plugins\cache\openai-primary-runtime\documents\26.813.12317\skills\documents\scripts")
sys.path.insert(0, str(TABLE_HELPER_DIR))

from generate_report import (  # noqa: E402
    BLUE,
    GRAY,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    add_body,
    add_direct_paragraph,
    add_h1,
    add_h2,
    add_page_number_section,
    add_sequence_field,
    configure_document,
    make_architecture_diagram,
    make_relationship_diagram,
    make_use_case_diagram,
    pfont,
    set_image_alt,
    set_run_font,
    set_table_geometry,
    shade_cell,
)
from table_geometry import apply_table_geometry  # noqa: E402


OUT_FILE = ROOT / "entregaveis" / "Relatorio_Evolucao_SAC_IFRR.docx"
ASSETS = Path(__file__).resolve().parent / "assets"


def add_field(paragraph, instruction: str, fallback: str, size: int = 12):
    """Insert a Word field that can be refreshed with Ctrl+A, F9."""
    run = paragraph.add_run()
    set_run_font(run, size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_table_caption(doc: Document, number: int, title: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run("Tabela ")
    set_run_font(run, 10)
    seq = paragraph.add_run()
    set_run_font(seq, 10)
    add_sequence_field(seq, "Tabela", number)
    run = paragraph.add_run(f" – {title}")
    set_run_font(run, 10)


def add_table(doc: Document, number: int, title: str, headers, rows, widths, source: str):
    add_table_caption(doc, number, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    apply_table_geometry(table, widths)
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "D9E5F3")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(heading)
        set_run_font(run, 9, bold=True)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.first_line_indent = Cm(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, 9)
    apply_table_geometry(table, widths)
    source_p = doc.add_paragraph(style="Source Note")
    run = source_p.add_run(source)
    set_run_font(run, 10)
    return table


def add_figure(doc: Document, image_path: Path, number: int, title: str, description: str, width_cm=15.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Cm(width_cm))
    set_image_alt(inline, f"Figura {number} – {title}", description)
    caption = doc.add_paragraph(style="Figure Caption")
    prefix = caption.add_run("Figura ")
    set_run_font(prefix, 10)
    seq = caption.add_run()
    set_run_font(seq, 10)
    add_sequence_field(seq, "Figura", number)
    suffix = caption.add_run(f" – {title}")
    set_run_font(suffix, 10)
    source = doc.add_paragraph(style="Source Note")
    run = source.add_run("Fonte: elaboração própria (2026).")
    set_run_font(run, 10)


def make_timeline() -> Path:
    """Create a compact historical timeline with only evidence-backed milestones."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    output = ASSETS / "figura_1_linha_do_tempo.png"
    image = Image.new("RGB", (1900, 950), "white")
    draw = ImageDraw.Draw(image)
    blue = "#294E82"
    light = "#EAF1F8"
    gray = "#626262"
    draw.text((520, 48), "Evolução técnica do Sistema de Alunos Cotistas", font=pfont(34, True), fill=blue)
    y = 470
    draw.line((145, y, 1755, y), fill=blue, width=6)
    points = [245, 685, 1120, 1560]
    labels = [
        ("Linha de base", "Portal local\nNode/Express/SQLite\nCRUD e risco inicial"),
        ("Reestruturação", "React + TypeScript\nAPI JSON e camadas\ncliente/servidor"),
        ("Integridade", "RBAC, bcrypt, sessões\ncontas transacionais\ne matrícula normalizada"),
        ("Análise histórica", "Cotas, CSV, gráficos\nsnapshots por período\ne exportação"),
    ]
    for x, (heading, text) in zip(points, labels):
        draw.ellipse((x - 30, y - 30, x + 30, y + 30), fill=light, outline=blue, width=5)
        left, top, right, bottom = x - 180, 185, x + 180, 395
        draw.rounded_rectangle((left, top, right, bottom), radius=18, fill=light, outline=blue, width=3)
        heading_font = pfont(23, True)
        body_font = pfont(19, False)
        bbox = draw.textbbox((0, 0), heading, font=heading_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, top + 24), heading, font=heading_font, fill="#111111")
        current_y = top + 75
        for line in text.split("\n"):
            bbox = draw.textbbox((0, 0), line, font=body_font)
            draw.text((x - (bbox[2] - bbox[0]) / 2, current_y), line, font=body_font, fill="#111111")
            current_y += 31
        draw.line((x, 395, x, y - 30), fill=gray, width=3)
    foot = (
        "A sequência é uma reconstrução técnica a partir do documento de origem, do relatório técnico anterior e da versão atual do código; "
        "o repositório não possui histórico Git suficiente para datar cada marco."
    )
    draw.multiline_text((185, 690), foot, font=pfont(20, False), fill=gray, spacing=10)
    image.save(output, "PNG", optimize=True)
    return output


def make_account_cycle_diagram() -> Path:
    """Show the current account lifecycle without exposing an operational password formula."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    output = ASSETS / "figura_5_ciclo_conta.png"
    image = Image.new("RGB", (1900, 980), "white")
    draw = ImageDraw.Draw(image)
    blue = "#294E82"
    light = "#EAF1F8"
    gray = "#666666"
    draw.text((520, 42), "Ciclo seguro da conta estudantil", font=pfont(34, True), fill=blue)

    nodes = [
        (85, 350, 385, 560, "Cadastro manual\nou CSV"),
        (475, 350, 775, 560, "Transação\naluno + conta"),
        (865, 350, 1165, 560, "Senha em hash\ne perfil alinhado"),
        (1255, 350, 1555, 560, "Login e\nsessão SQLite"),
    ]
    for x0, y0, x1, y1, text in nodes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=20, fill=light, outline=blue, width=3)
        font = pfont(25, True)
        lines = text.split("\n")
        yy = y0 + 64
        for line in lines:
            box = draw.textbbox((0, 0), line, font=font)
            draw.text(((x0 + x1 - (box[2] - box[0])) / 2, yy), line, font=font, fill="#111111")
            yy += 40
    for x in (385, 775, 1165):
        draw.line((x + 12, 455, x + 78, 455), fill=blue, width=4)
        draw.polygon([(x + 78, 455), (x + 62, 447), (x + 62, 463)], fill=blue)

    draw.rounded_rectangle((520, 690, 1375, 840), radius=20, fill="#F2F2F2", outline=gray, width=3)
    note = "Na inicialização e antes de novo cadastro: limpar conta de papel Aluno sem matrícula ativa, remover sessões associadas e preservar contas da equipe e snapshots históricos."
    body = pfont(22, False)
    words = note.split()
    lines, current = [], ""
    for word in words:
        proposal = f"{current} {word}".strip()
        if draw.textbbox((0, 0), proposal, font=body)[2] > 780:
            lines.append(current)
            current = word
        else:
            current = proposal
    if current:
        lines.append(current)
    yy = 730
    for line in lines:
        box = draw.textbbox((0, 0), line, font=body)
        draw.text(((520 + 1375 - (box[2] - box[0])) / 2, yy), line, font=body, fill="#111111")
        yy += 30
    draw.line((1405, 560, 1265, 690), fill=gray, width=3)
    draw.text((170, 885), "A senha inicial é tratada como procedimento operacional de primeiro acesso e não é exposta no relatório.", font=pfont(19, False), fill=gray)
    image.save(output, "PNG", optimize=True)
    return output


def add_cover(doc: Document):
    for line in [
        "MINISTÉRIO DA EDUCAÇÃO",
        "SECRETARIA DE EDUCAÇÃO PROFISSIONAL E TECNOLÓGICA",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA",
        "[CAMPUS / DIRETORIA]",
        "[CURSO]",
    ]:
        add_direct_paragraph(doc, line, size=14, bold=True, after=0)
    add_direct_paragraph(doc, "", before=70)
    add_direct_paragraph(doc, "[NOME DO(A) AUTOR(A)]", size=16, bold=True)
    add_direct_paragraph(doc, "", before=90)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO DE EVOLUÇÃO", size=18, bold=True)
    add_direct_paragraph(doc, "DO SISTEMA DE ALUNOS COTISTAS (SAC)", size=18, bold=True)
    add_direct_paragraph(doc, "", before=75)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO/CIENTÍFICO", size=14, bold=True)
    add_direct_paragraph(doc, "", before=75)
    add_direct_paragraph(doc, "[CIDADE – UF]", size=14, bold=True)
    add_direct_paragraph(doc, "2026", size=14, bold=True)


def add_title_page(doc: Document):
    doc.add_page_break()
    add_direct_paragraph(doc, "[NOME DO(A) AUTOR(A)]", size=14, bold=True)
    add_direct_paragraph(doc, "", before=85)
    add_direct_paragraph(doc, "RELATÓRIO TÉCNICO DE EVOLUÇÃO", size=16, bold=True)
    add_direct_paragraph(doc, "DO SISTEMA DE ALUNOS COTISTAS (SAC)", size=16, bold=True)
    add_direct_paragraph(doc, "", before=60)
    nature = (
        "Relatório técnico/científico elaborado como documento de apoio ao Trabalho de Conclusão de Curso, "
        "apresentado ao [CURSO] do Instituto Federal de Educação, Ciência e Tecnologia de Roraima, "
        "como requisito parcial para obtenção do título de [GRAU].\n\n"
        "Orientador(a): [NOME DO(A) ORIENTADOR(A)]"
    )
    add_direct_paragraph(doc, nature, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    add_direct_paragraph(doc, "", before=115)
    add_direct_paragraph(doc, "[CIDADE – UF]", size=14)
    add_direct_paragraph(doc, "2026", size=14)


def add_front_matter(doc: Document):
    doc.add_page_break()
    add_direct_paragraph(doc, "RESUMO", size=14, bold=True, after=24)
    abstract = (
        "Este relatório técnico/científico consolida a evolução do Sistema de Alunos Cotistas (SAC) a partir de um "
        "documento de linha de base e da versão atual do projeto. A análise preserva o contexto institucional de "
        "permanência e acompanhamento de estudantes cotistas, separando-o das funcionalidades comprovadas no "
        "código atual. A evolução abrange a substituição da interface renderizada no servidor por uma aplicação React "
        "com TypeScript, a adoção de uma API JSON em Express, a organização de persistência em SQLite e o reforço "
        "da autenticação, das sessões e do controle de acesso por papéis. O relatório apresenta tabelas comparativas, "
        "arquitetura, casos de uso, relacionamentos lógicos, ciclo de contas, tratamento de notas e faltas, importação "
        "CSV, estatísticas, snapshots por período e exportação de informações. A evidência técnica considerada inclui "
        "o relatório anexado, o relatório técnico anterior, a inspeção do repositório e verificações de build, lint, "
        "sintaxe e testes focalizados. O resultado é uma visão organizada da transição de um portal acadêmico local "
        "para uma solução web com responsabilidades separadas, maior integridade de identidade estudantil e apoio "
        "a análises acadêmicas por escopo de acesso."
    )
    add_direct_paragraph(doc, abstract, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    add_direct_paragraph(doc, "Palavras-chave: estudantes cotistas; evolução de software; React; autenticação; estatísticas acadêmicas.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0, before=12)

    doc.add_page_break()
    add_direct_paragraph(doc, "ABSTRACT", size=14, bold=True, after=24)
    english = (
        "This technical/scientific report consolidates the evolution of the Sistema de Alunos Cotistas (SAC) from "
        "a baseline document and the current project version. It preserves the institutional context of student retention "
        "while distinguishing it from features supported by the current source code. The evolution includes the move from "
        "server-rendered pages to a React and TypeScript application, a JSON API in Express, SQLite persistence, and "
        "stronger authentication, sessions and role-based access control. Comparative tables, architecture, use cases, "
        "logical relationships, account lifecycle, academic indicators, CSV import, statistics, period snapshots and exports "
        "are presented. The resulting system has clearer responsibility boundaries, improved student identity integrity and "
        "role-scoped academic analysis."
    )
    add_direct_paragraph(doc, english, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0)
    add_direct_paragraph(doc, "Keywords: quota students; software evolution; React; authentication; academic statistics.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.0, first_indent=0, before=12)

    doc.add_page_break()
    add_direct_paragraph(doc, "LISTA DE FIGURAS", size=14, bold=True, after=20)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, ' TOC \\h \\z \\c "Figura" ', "Atualize a lista de figuras no Microsoft Word, se necessário.")

    doc.add_page_break()
    add_direct_paragraph(doc, "LISTA DE TABELAS", size=14, bold=True, after=20)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, ' TOC \\h \\z \\c "Tabela" ', "Atualize a lista de tabelas no Microsoft Word, se necessário.")

    doc.add_page_break()
    add_direct_paragraph(doc, "LISTA DE SIGLAS", size=14, bold=True, after=20)
    siglas = [
        ("API", "Application Programming Interface"),
        ("CSV", "Comma-Separated Values"),
        ("IFRR", "Instituto Federal de Educação, Ciência e Tecnologia de Roraima"),
        ("RBAC", "Role-Based Access Control"),
        ("SAC", "Sistema de Alunos Cotistas"),
        ("SPA", "Single Page Application"),
        ("TCC", "Trabalho de Conclusão de Curso"),
    ]
    for sigla, meaning in siglas:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(f"{sigla} – ")
        set_run_font(run, 12, bold=True)
        run = p.add_run(meaning)
        set_run_font(run, 12)

    doc.add_page_break()
    add_direct_paragraph(doc, "SUMÁRIO", size=14, bold=True, after=20)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ', "Atualize o sumário no Microsoft Word, se necessário.")


def add_introduction(doc: Document):
    add_h1(doc, "1 INTRODUÇÃO", page_break=False)
    add_body(doc, "O acompanhamento acadêmico de estudantes cotistas demanda informações de matrícula, curso, turma, disciplina, modalidade de cota, notas, faltas e registros de apoio. O relatório anexado ao presente trabalho apresenta esse problema no contexto de permanência, evasão e acompanhamento institucional, descrevendo uma primeira solução de portal acadêmico local. A versão atual do Sistema de Alunos Cotistas (SAC) amplia essa base e transforma o produto em uma aplicação web com interface separada, API, controle de acesso e recursos de análise por período.")
    add_body(doc, "Este documento une as partes comuns entre o relatório de origem e o relatório técnico anteriormente produzido, organizando-as em uma narrativa de evolução. Para não transformar intenção em fato, elementos do documento inicial são identificados como linha de base ou recomendação histórica. Funcionalidades da versão atual são descritas somente quando encontradas na implementação e nas verificações executadas.")
    add_h2(doc, "1.1 Objetivo geral")
    add_body(doc, "Apresentar a evolução técnica e funcional do SAC desde a primeira versão documentada até a versão atual, evidenciando mudanças de arquitetura, segurança, integridade de dados, acompanhamento acadêmico e experiência de uso.")
    add_h2(doc, "1.2 Objetivos específicos")
    add_body(doc, "Os objetivos específicos são: identificar o que foi preservado da primeira proposta; comparar as tecnologias e as responsabilidades das duas fases; descrever os mecanismos atuais de autenticação, autorização e ciclo de contas; registrar a evolução de cadastro, cotas, risco, estatísticas e históricos; e indicar limitações e próximos passos para uso institucional.")
    add_h2(doc, "1.3 Delimitação e transparência de evidência")
    add_body(doc, "O relatório de origem reúne um portal acadêmico, trechos de documentação do protótipo e um contexto institucional de acompanhamento de cotistas. Ele contém também recomendações de integração e governança. Por isso, integração com SUAP/CTI, autenticação multifator, infraestrutura de rede, criptografia em repouso e conformidade jurídica integral são tratados neste texto como possibilidades futuras, e não como recursos implementados. Como o repositório atual possui histórico Git reduzido, a cronologia é uma reconstrução técnica baseada nos documentos e no código, não uma auditoria de datas de desenvolvimento.")


def add_baseline(doc: Document):
    add_h1(doc, "2 LINHA DE BASE E ELEMENTOS CONSOLIDADOS")
    add_body(doc, "A primeira parte documentada descreve um Portal Acadêmico Integrado executado localmente com Node.js, Express e SQLite. A interface era renderizada no servidor, com referências a EJS/Jinja2 e Bootstrap. Já estavam presentes operações de cadastro, consulta, busca, exclusão, formatação de telefone e uma regra inicial de atenção por notas e faltas. Essa fase também registrou o interesse institucional em correlacionar desempenho, frequência e modalidades de cota.")
    add_body(doc, "A versão atual não contém EJS, Jinja2, Bootstrap, views ou notebooks Python no seu código. Essas tecnologias pertencem exclusivamente à linha de base. Permanecem como fundamentos tecnológicos o Node.js, o Express e o SQLite, mas agora ocupam responsabilidades específicas de API e persistência.")
    add_table(
        doc, 1, "Fontes usadas na reconstrução da evolução",
        ["Fonte", "Papel no relatório", "Tratamento adotado"],
        [
            ("Relatório Técnico IFRR anexado", "Linha de base, contexto de cotas e diagnóstico histórico.", "Preserva fatos históricos; recomendações não são apresentadas como entrega atual."),
            ("Relatório técnico anterior do SAC", "Estrutura acadêmica, arquitetura atual e diagramas.", "Atualizado com a evolução de contas, históricos e filtros."),
            ("Código-fonte e verificações atuais", "Evidência de funcionalidades implementadas.", "Usado para descrever a versão vigente e seus limites."),
        ],
        [2000, 3200, 3840],
        "Fonte: elaboração própria, com base nos documentos fornecidos e no repositório do SAC (2026).",
    )
    add_h2(doc, "2.1 Partes comuns preservadas")
    add_body(doc, "A motivação de apoiar a permanência acadêmica dos estudantes cotistas foi preservada. Também foram mantidos o cadastro de estudantes, o uso de notas e faltas, a indicação de situação que requer atenção e o uso de tecnologias web locais de fácil operação. O que mudou foi a profundidade da modelagem, a proteção dos dados e a capacidade de analisar a evolução por perfis e períodos.")
    add_table(
        doc, 2, "Elementos comuns preservados e expandidos",
        ["Elemento comum", "Fase inicial", "Versão atual"],
        [
            ("Cadastro acadêmico", "Registro individual e operações básicas.", "Cadastro manual e CSV, validações, até 20 matérias e conta de aluno transacional."),
            ("Notas e faltas", "Indicador simples de risco por desempenho e frequência.", "Indicadores por matéria e agregados; situação Regular, Alerta ou Em Risco."),
            ("Cotas", "Contexto de modalidades e acompanhamento institucional.", "Lista centralizada no cadastro e filtros administrativos de estatísticas."),
            ("Acesso à informação", "Portal local com consulta e administração inicial.", "Perfis Aluno, Professor, Diretor e Admin com escopo verificado pela API."),
            ("Tomada de decisão", "Intenção de acompanhar permanência e evasão.", "Gráficos, exportação, impressão e snapshots comparáveis por período."),
        ],
        [1700, 3650, 3690],
        "Fonte: elaboração própria (2026).",
    )


def add_evolution(doc: Document, figures):
    add_h1(doc, "3 TRAJETÓRIA DE EVOLUÇÃO DO PROJETO")
    add_body(doc, "A evolução foi organizada por marcos de engenharia, e não por datas exatas. A Figura 1 sintetiza a passagem do portal de linha de base para a solução atual. O primeiro marco manteve uma solução local funcional; o segundo separou cliente e servidor; o terceiro fortaleceu identidade, sessão e escopo; e o quarto agregou histórico e análise dinâmica.")
    add_figure(doc, figures["timeline"], 1, "Linha do tempo técnica do SAC", "Marcos técnicos: linha de base, reestruturação React/API, integridade de identidade e análise histórica.")
    add_table(
        doc, 3, "Matriz de evolução arquitetural e tecnológica",
        ["Dimensão", "Linha de base", "Estado atual", "Ganho de engenharia"],
        [
            ("Interface", "Páginas renderizadas no servidor; referências a EJS/Jinja2 e Bootstrap.", "React 19, TypeScript, Vite, CSS próprio, páginas e componentes.", "Interface modular, contratos de tipos e experiência de SPA."),
            ("Comunicação", "Rotas HTTP acopladas à apresentação.", "API JSON sob /api consumida pelo cliente React.", "Fronteira clara entre tela e regras de negócio."),
            ("Servidor", "Express concentrando fluxo e acesso ao banco.", "src/server.js, rotas de autenticação e domínio, middleware e respostas padronizadas.", "Responsabilidades mais fáceis de manter e testar."),
            ("Persistência", "Dados acadêmicos básicos em alunos.", "alunos, usuarios, sessoes, historico_academico e movimentacoes.", "Contas, sessão e snapshots passam a ter representação própria."),
            ("Operação", "Inicialização manual.", "quickstart compila React, inicia Express e abre o navegador.", "Execução local mais reproduzível."),
        ],
        [1300, 2600, 2700, 2440],
        "Fonte: elaboração própria, com base no relatório de origem e no código atual do SAC (2026).",
    )
    add_h2(doc, "3.1 Incidentes e correções que orientaram a evolução")
    add_body(doc, "A documentação inicial registra uma falha de escopo/exportação em que uma função de hash não estava disponível no módulo esperado. Durante a evolução, problemas de identidade de aluno também foram analisados: uma conta antiga pode existir com perfil desatualizado, e uma matrícula reaproveitada não pode herdar uma conta órfã. Os mecanismos atuais tratam esses casos de forma centralizada e evitando substituição automática de senha personalizada.")
    add_table(
        doc, 4, "Problemas históricos e tratamento atual",
        ["Situação", "Risco", "Tratamento implementado", "Evidência técnica"],
        [
            ("Função de hash indisponível no protótipo", "Login e inicialização falhavam por escopo inadequado.", "Utilitários organizados e exportados no nível de módulo.", "src/lib/utils.js"),
            ("Conta de aluno divergente", "Matrícula podia apontar para perfil antigo.", "Backfill e reconciliação de perfil; senha customizada preservada.", "src/lib/database.js"),
            ("Matrícula reutilizada", "Conta órfã poderia bloquear ou transferir identidade.", "Limpeza seletiva no início e antes de novo cadastro.", "cleanupOrphanStudentAccounts()"),
            ("Permissão genérica para aluno", "Identidade estudantil poderia ser convertida indevidamente.", "Tela e API de equipe recusam criação/alteração genérica de Aluno.", "src/routes/api.js; PermissionsPage.tsx"),
            ("Histórico contaminado", "Aluno cadastrado depois alteraria período fechado.", "Snapshot usa elenco capturado, com sobrescrita explícita.", "historico_academico"),
        ],
        [1700, 2450, 3000, 1890],
        "Fonte: elaboração própria, com base no relatório de origem e no código atual (2026).",
    )


def add_architecture(doc: Document, figures):
    add_h1(doc, "4 ARQUITETURA E ESTRUTURA ATUAL DO PROJETO")
    add_body(doc, "A arquitetura atual utiliza uma aplicação React de página única como cliente e um servidor Express como API e distribuidor dos arquivos estáticos compilados. O banco SQLite concentra os dados acadêmicos, as contas, as sessões e os retratos históricos. A interface não é considerada barreira suficiente para dados sensíveis: o servidor revalida perfil, matrícula, curso, disciplina e turma antes de responder ou alterar registros.")
    add_figure(doc, figures["architecture"], 2, "Arquitetura em camadas do SAC", "Fluxo entre navegador, cliente React, servidor Express e banco SQLite, com responsabilidades de cada camada.")
    add_table(
        doc, 5, "Estrutura atual do projeto e responsabilidades",
        ["Camada", "Principais arquivos ou diretórios", "Responsabilidade"],
        [
            ("Inicialização", "src/server.js; src/scripts/quickstart.js", "Configura Express, entrega a SPA compilada e simplifica a execução local."),
            ("Rotas HTTP", "src/routes/auth.js; src/routes/api.js", "Autenticação, alunos, usuários de equipe, CSV, estatísticas, períodos e backup."),
            ("Domínio e dados", "src/lib/database.js; auth.js; permissions.js; utils.js", "SQL, regras acadêmicas, autenticação, escopo e normalização."),
            ("Sessões", "src/lib/sessionStore.js", "Sessões persistentes em SQLite e expiração."),
            ("Cliente web", "teste-react/src/pages; components; utils; types", "Telas, tabela de alunos, formulários, gráficos, exportação e contratos TypeScript."),
            ("Operação", "QUICKSTART.md; package.json", "Comandos locais, build e primeira criação administrativa."),
        ],
        [1600, 3550, 3890],
        "Fonte: elaboração própria, com base na estrutura do repositório do SAC (2026).",
    )
    add_h2(doc, "4.1 Modelo lógico de dados")
    add_body(doc, "O modelo atual é composto por alunos, usuários, sessões, histórico acadêmico e movimentações. Há relações lógicas deliberadas em vez de chaves estrangeiras físicas em todos os vínculos. Isso permite que um snapshot acadêmico continue disponível mesmo se o cadastro ativo do estudante for removido. A tabela movimentacoes está prevista para auditoria, mas ainda não está integrada como trilha consultável do fluxo principal.")
    add_figure(doc, figures["relationships"], 3, "Relacionamentos lógicos de dados do SAC", "Tabelas alunos, usuários, histórico acadêmico, sessões e movimentações, com observações sobre retenção e vínculos lógicos.")


def add_functional_evolution(doc: Document):
    add_h1(doc, "5 EVOLUÇÃO FUNCIONAL: ACOMPANHAMENTO, COTAS E ESTATÍSTICAS")
    add_body(doc, "O SAC atual mantém a proposta de acompanhamento de estudantes cotistas e a estende com dados por matéria, filtros, gráficos e comparação temporal. As regras de risco preexistentes são preservadas como referência: o sistema trabalha com situação Regular, Alerta ou Em Risco usando notas e frequência. Como parâmetros pedagógicos precisam de validação institucional, este relatório descreve a regra implementada sem declará-la como política oficial do IFRR.")
    add_table(
        doc, 6, "Evolução das funcionalidades acadêmicas",
        ["Domínio", "Fase inicial", "Versão atual"],
        [
            ("Gestão de estudantes", "CRUD, busca e dados básicos.", "Cadastro manual, edição, descrição, CSV em lote, modelo de arquivo e validações."),
            ("Matérias", "Indicadores gerais por aluno.", "Até 20 matérias por aluno, notas, faltas, faltas justificadas, aulas e situação por componente."),
            ("Risco", "Nota abaixo de 70 ou falta acima de 25%, conforme documento de origem.", "Indicadores agregados e por matéria; situação Em Risco, Alerta e Regular; resumo legado preservado."),
            ("Cotas", "Mapeamento de modalidades e análise desejada.", "Opções centralizadas no cadastro e filtros administrativos por modalidade de cota."),
            ("Descrição", "Não descrita como recurso de interação.", "Prévia acessível no hover/foco e modal completo acionado pelo nome do aluno."),
            ("Períodos", "Sem histórico persistido comparável.", "Quatro bimestres e dois semestres com snapshots explícitos e sobrescrita confirmada."),
            ("Estatísticas", "Intenção de correlacionar desempenho e frequência.", "Dados dinâmicos por escopo, cinco gráficos SVG, CSV, PDF e impressão."),
        ],
        [1600, 3320, 4120],
        "Fonte: elaboração própria, com base no relatório de origem e na implementação atual (2026).",
    )
    add_h2(doc, "5.1 Cotas e filtros")
    add_body(doc, "A lista de modalidades de cota foi centralizada no cliente para que o cadastro e as estatísticas usem a mesma referência. Os códigos presentes no projeto incluem AC, PCD_AC, L1, L2, L5, L6, L9, L10, L13 e L14. Essa lista representa as opções atualmente cadastradas no sistema; sua aderência a editais ou normas institucionais vigentes deve ser confirmada pela área responsável antes de uma implantação oficial.")
    add_h2(doc, "5.2 Estatísticas com escopo")
    add_body(doc, "O estudante visualiza apenas o próprio desempenho; o professor visualiza a turma vinculada; a direção visualiza o curso atribuído; e a administração tem visão ampla e filtros adicionais de curso, turma e cota. Os cinco formatos disponíveis são pizza, colunas, linhas, barras e área. A exportação em CSV, a geração de PDF e a impressão permitem levar a análise para reuniões ou acompanhamento externo sem alterar a fonte de dados.")
    add_h2(doc, "5.3 Históricos por período")
    add_body(doc, "O snapshot é uma fotografia salva por administrador. Após o salvamento, o período usa exclusivamente o conjunto de estudantes capturado, impedindo que cadastros posteriores alterem o resultado de um período fechado. Também é possível recuperar o retrato de estudante removido do cadastro atual. O administrador precisa confirmar a sobrescrita, pois ela altera explicitamente a fotografia existente.")


def add_security(doc: Document, figures):
    add_h1(doc, "6 AUTENTICAÇÃO, AUTORIZAÇÃO E INTEGRIDADE DE CONTAS")
    add_body(doc, "A evolução de segurança parte da necessidade de corrigir um problema de escopo no protótipo e de tornar o fluxo de alunos previsível. A versão atual utiliza hash bcrypt, sessões persistentes em SQLite, cookies HTTP-only, SameSite=Lax e atributo Secure em produção. Mudanças de senha incrementam uma versão de sessão, tornando sessões anteriores inválidas. Alterações de estado com origem cruzada não esperada são bloqueadas.")
    add_table(
        doc, 7, "Evolução dos controles de identidade e acesso",
        ["Controle", "Tratamento atual", "Finalidade"],
        [
            ("Matrícula", "Normalização por remoção de espaços externos e caixa padronizada; índice único sem diferenciar maiúsculas/minúsculas.", "Evitar duplicidade e ambiguidade de login."),
            ("Conta de aluno", "Criada junto ao aluno, manualmente ou via CSV, dentro da mesma transação.", "Evitar estudante cadastrado sem acesso correspondente."),
            ("Login", "Usuário, matrícula ou nome completo único, quando permitido pelo perfil.", "Compatibilidade controlada sem escolher usuários ambíguos."),
            ("Senha", "bcrypt, validação de tamanho e troca autenticada; fórmula operacional não exposta no relatório.", "Reduzir exposição de credenciais."),
            ("Sessão", "SQLite, expiração, versão de sessão e cookie protegido.", "Manter autenticação persistente com invalidação controlada."),
            ("Tentativas", "Limite de 5 falhas em 15 minutos por IP e login normalizado, em memória.", "Reduzir tentativa repetida sem bloquear toda a rede."),
            ("RBAC", "Admin, Diretor, Professor e Aluno com escopo validado na API.", "Restringir dados por matrícula, turma, disciplina e curso."),
        ],
        [1600, 4700, 2740],
        "Fonte: elaboração própria, com base na implementação do SAC e na OWASP (2026).",
    )
    add_h2(doc, "6.1 Ciclo de conta estudantil")
    add_body(doc, "O ciclo atual evita dois problemas recorrentes: a existência de aluno sem conta e a reutilização de matrícula vinculada a uma conta antiga. Na inicialização ocorre um backfill idempotente para contas faltantes e uma reconciliação de perfil. Quando o perfil diverge, dados cadastrais podem ser alinhados sem substituir senha personalizada; somente uma credencial comprovadamente ainda associada ao procedimento inicial anterior pode ser redefinida em correção controlada.")
    add_figure(doc, figures["account"], 4, "Ciclo seguro de criação e limpeza de contas estudantis", "Fluxo de cadastro, criação transacional, hash, sessão e limpeza de conta órfã sem expor senha operacional.")
    add_body(doc, "Uma conta de papel Aluno sem matrícula ativa é considerada órfã e é removida juntamente com suas sessões, preservando contas da equipe e snapshots acadêmicos. O sistema também possui gatilho para revogar a conta ao excluir o cadastro do aluno. A gestão genérica de permissões não cria nem converte contas estudantis, impedindo que uma senha de aluno seja transferida para outra identidade ou para papel de equipe.")


def add_use_cases_and_validation(doc: Document, figures):
    add_h1(doc, "7 CASOS DE USO, VALIDAÇÃO E QUALIDADE")
    add_h2(doc, "7.1 Casos de uso")
    add_body(doc, "A Figura 5 representa os atores da versão atual. Ações como autenticar-se, encerrar sessão, alterar senha e consultar estatísticas existem em níveis compatíveis com cada papel. O que cada perfil pode visualizar, importar, administrar ou salvar é definido no servidor. A figura mostra a visão funcional; os filtros e a validação de escopo fazem parte do contrato da API.")
    add_figure(doc, figures["use_case"], 5, "Casos de uso do Sistema de Alunos Cotistas", "Atores Aluno, Professor, Diretor e Administrador com operações de acordo com o respectivo escopo.")
    add_h2(doc, "7.2 Verificações realizadas")
    add_body(doc, "A validação foi executada em proporção ao risco das alterações. Além de compilar o cliente, foram verificados cenários de ciclo de contas, cadastro manual, importação CSV, login por matrícula e nome único, colisões de matrícula, escopo de perfis e snapshots. Testes isolados evitaram modificar a base de dados principal durante a reprodução de problemas.")
    add_table(
        doc, 8, "Evidências de validação da versão atual",
        ["Verificação", "Comando ou cenário", "Resultado e limite"],
        [
            ("Testes de conta", "node --test test/student-account-cleanup.test.js", "Três testes aprovados para limpeza, sessões e reutilização de matrícula."),
            ("Cliente", "npm run build", "Build TypeScript/Vite concluído sem erro."),
            ("Lint", "npm --prefix teste-react run lint", "Análise estática do cliente concluída sem erro."),
            ("Servidor", "node --check em módulos principais", "Sintaxe validada sem erro."),
            ("Integração focalizada", "Bases SQLite temporárias e HTTP isolado", "Fluxos sensíveis de conta e histórico reproduzidos e verificados."),
            ("Dependências", "npm audit --omit=dev", "Sem vulnerabilidades reportadas no momento da verificação."),
        ],
        [1600, 3400, 4040],
        "Fonte: elaboração própria (2026).",
    )
    add_h2(doc, "7.3 Interpretação da evidência")
    add_body(doc, "Os resultados confirmam os fluxos mais críticos abordados nesta evolução, sobretudo a identidade de estudantes, o tratamento de contas órfãs e a fidelidade de snapshots. A cobertura automatizada ainda é restrita: não há suíte abrangente de interface, API, RBAC e importação. Logo, os resultados não equivalem a certificação de qualidade total nem substituem homologação com usuários e responsáveis acadêmicos.")


def add_limitations(doc: Document):
    add_h1(doc, "8 LIMITAÇÕES, RECOMENDAÇÕES E CONTINUIDADE INSTITUCIONAL")
    add_body(doc, "A versão atual concentra funcionalidades importantes para operação local e demonstra evolução de engenharia, mas ainda requer preparação para uso institucional ampliado. As recomendações abaixo distinguem melhorias diretamente observadas no repositório de integrações ou políticas mencionadas apenas como intenção no relatório de origem.")
    add_table(
        doc, 9, "Plano de melhorias e prioridades",
        ["Prioridade", "Melhoria", "Motivo"],
        [
            ("Alta", "Completar README com arquitetura, variáveis, perfis, importação e manutenção.", "O QUICKSTART existe, mas o README principal ainda é insuficiente para transferência de conhecimento."),
            ("Alta", "Configurar HTTPS, SESSION_SECRET forte e backup externo com teste de restauração.", "Controles de infraestrutura não são resolvidos apenas pelo código local."),
            ("Alta", "Ampliar testes de API, RBAC, login, CSV e interface.", "A suíte automatizada atual é focalizada no ciclo de contas."),
            ("Média", "Implementar rotação de backup e retenção configurável.", "A cópia atual substitui um único arquivo de backup; não há retenção de 20 versões."),
            ("Média", "Integrar movimentacoes como trilha de auditoria consultável.", "A tabela existe, porém não é gravada/lida pelo fluxo principal."),
            ("Média", "Homologar parâmetros de risco e categorias de cota com área pedagógica.", "Regra de nota/frequência e modalidades devem ser formalizadas institucionalmente."),
            ("Evolutiva", "Planejar integração com CTI/SUAP, banco mais escalável, MFA e recuperação assistida.", "São direcionamentos futuros do contexto institucional, não recursos implementados."),
        ],
        [1200, 4000, 3840],
        "Fonte: elaboração própria, a partir da análise técnica e de recomendações históricas do documento de origem (2026).",
    )
    add_h2(doc, "8.1 Cuidados de produção")
    add_body(doc, "O rate limit atual é mantido em memória, portanto não é compartilhado entre múltiplos processos e é reiniciado com o servidor. A persistência SQLite é adequada ao escopo local, mas uma implantação multiusuário ampliada deve avaliar concorrência, cópias de segurança, monitoramento e uma estratégia de banco gerenciado. Também não há chaves estrangeiras físicas para todos os vínculos; a regra de retenção de snapshots precisa ser preservada caso o modelo seja migrado.")


def add_conclusion(doc: Document):
    add_h1(doc, "9 CONSIDERAÇÕES FINAIS")
    add_body(doc, "O SAC evoluiu de uma proposta de portal acadêmico local, com cadastro e indicadores iniciais de risco, para uma aplicação web organizada em cliente React, API Express e persistência SQLite. A motivação de acompanhar estudantes cotistas permanece, mas foi materializada em uma arquitetura com escopo por papel, importação, estatísticas, exportação e retratos históricos por período.")
    add_body(doc, "Os ganhos centrais estão na separação de responsabilidades, na criação transacional de aluno e conta, na normalização de matrícula, na reconciliação segura de perfis, na limpeza de contas órfãs e na proteção dos dados históricos contra cadastros posteriores. A versão atual oferece uma base consistente para continuidade acadêmica e técnica, desde que as regras pedagógicas, as modalidades de cota e os requisitos de produção sejam homologados pelas áreas responsáveis.")
    add_body(doc, "Para entrega como TCC, o autor deve preencher os campos de identificação, atualizar os campos de sumário e listas no Word, confirmar exigências específicas do curso e submeter a redação à revisão do orientador. O documento foi estruturado conforme o Manual do IFRR como referência formal, sem pretender substituir as normas ABNT e o PPC aplicáveis à banca.")


def add_references(doc: Document):
    add_h1(doc, "REFERÊNCIAS")
    entries = [
        "EXPRESSJS. Production best practices: security. [S. l.], 2026. Disponível em: <https://expressjs.com/en/advanced/best-practice-security/>. Acesso em: 17 ago. 2026.",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA. Manual de normas para elaboração de trabalhos acadêmicos do IFRR. Boa Vista: IFRR, [s. d.]. Disponível em: <https://boavista.ifrr.edu.br/aluno/manual-detrabalhos-academicos-do-ifrr/manual-de-trabalhos-academicos-do-ifrr/at_download/file>. Acesso em: 17 ago. 2026.",
        "INSTITUTO FEDERAL DE EDUCAÇÃO, CIÊNCIA E TECNOLOGIA DE RORAIMA. Resolução CONSUP/IFRR nº 730, de 30 de março de 2023. Estabelece normas e diretrizes para a elaboração do Trabalho de Conclusão de Curso (TCC) dos cursos de graduação no âmbito do IFRR. Boa Vista: IFRR, 2023. Disponível em: <https://www.ifrr.edu.br/documents/1566/Resolu%C3%A7%C3%A3o_n._730_-2023.pdf>. Acesso em: 17 ago. 2026.",
        "OPEN WEB APPLICATION SECURITY PROJECT. Authentication cheat sheet. [S. l.], 2026. Disponível em: <https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html>. Acesso em: 17 ago. 2026.",
        "RELATÓRIO TÉCNICO IFRR. Documento de referência fornecido pelo autor. [S. l.: s. n.], s. d.",
        "REACT. React documentation. [S. l.], 2026. Disponível em: <https://react.dev/>. Acesso em: 17 ago. 2026.",
        "SISTEMA DE ALUNOS COTISTAS. Código-fonte da aplicação: versão analisada. [S. l.]: repositório do projeto, 2026.",
        "SISTEMA DE ALUNOS COTISTAS. Relatório técnico de desenvolvimento e evolução do SAC. [S. l.: s. n.], 2026.",
        "SQLITE. Transaction. [S. l.], 2026. Disponível em: <https://www.sqlite.org/lang_transaction.html>. Acesso em: 17 ago. 2026.",
    ]
    for entry in entries:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(entry)
        set_run_font(run, 12)


def add_appendix(doc: Document):
    add_h1(doc, "APÊNDICE A – MATRIZ DE EVIDÊNCIAS DA VERSÃO ATUAL")
    add_body(doc, "A matriz relaciona recursos descritos neste relatório a componentes do projeto. Ela não substitui a leitura do código, mas facilita manutenção, banca e continuidade do trabalho.")
    add_table(
        doc, 10, "Matriz de funcionalidades e componentes",
        ["Funcionalidade", "Evidência principal", "Componente"],
        [
            ("Login e sessões", "Rotas de login, logout, senha e sessão persistente.", "src/routes/auth.js; src/lib/auth.js; src/lib/sessionStore.js"),
            ("RBAC", "Escopo por matrícula, curso, disciplina e turma.", "src/lib/permissions.js; src/routes/api.js"),
            ("Cadastro e conta", "Transação, normalização, backfill e limpeza de órfãs.", "src/lib/database.js"),
            ("Importação CSV", "Validação, limites, modelo e resultado por linha.", "src/routes/api.js; RegistrationPage.tsx"),
            ("Estatísticas", "Agregação, filtros, gráficos SVG e exportação.", "StatisticsPage.tsx; statisticsPdf.ts"),
            ("Histórico", "Snapshot por período e evolução anual.", "src/lib/database.js; src/routes/api.js"),
            ("Descrição de aluno", "Tooltip, link e modal acessível.", "components/StudentTable.tsx"),
            ("Operação local", "Build, servidor e abertura no navegador.", "src/scripts/quickstart.js; QUICKSTART.md"),
        ],
        [1800, 3500, 3740],
        "Fonte: elaboração própria, com base no repositório do SAC (2026).",
    )
    add_h1(doc, "APÊNDICE B – ROTEIRO DE REVISÃO ANTES DA BANCA")
    add_body(doc, "Antes de submeter este arquivo, preencher nome, curso, campus, cidade, grau e orientação; atualizar sumário, lista de figuras e lista de tabelas no Microsoft Word; confirmar a modalidade de TCC e as exigências do PPC; revisar a aderência das cotas e dos parâmetros de risco; e remover ou complementar qualquer informação que a banca exija em formato específico.")
    add_table(
        doc, 11, "Checklist de entrega técnica",
        ["Etapa", "Ação", "Resultado esperado"],
        [
            ("1", "Preencher campos da capa e folha de rosto.", "Identificação acadêmica completa."),
            ("2", "Atualizar campos do Word com Ctrl+A e F9.", "Sumário e listas com paginação correta."),
            ("3", "Executar npm test, npm run build e lint do cliente.", "Evidências técnicas atualizadas."),
            ("4", "Revisar perfis e escopos em homologação.", "Aluno, Professor, Diretor e Admin limitados corretamente."),
            ("5", "Confirmar regras pedagógicas e cotas.", "Aderência à decisão institucional e ao edital aplicável."),
            ("6", "Configurar produção com TLS, segredo e backup.", "Ambiente preparado para uso além do contexto local."),
        ],
        [800, 4300, 3940],
        "Fonte: elaboração própria (2026).",
    )


def build():
    figures = {
        "timeline": make_timeline(),
        "architecture": make_architecture_diagram(),
        "relationships": make_relationship_diagram(),
        "account": make_account_cycle_diagram(),
        "use_case": make_use_case_diagram(),
    }
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Relatório Técnico de Evolução do Sistema de Alunos Cotistas (SAC)"
    doc.core_properties.subject = "Documento de apoio a TCC"
    doc.core_properties.author = "Sistema de Alunos Cotistas"
    doc.core_properties.comments = ""
    doc.core_properties.keywords = "SAC, IFRR, TCC, evolução de software, estudantes cotistas"
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)

    add_cover(doc)
    add_title_page(doc)
    add_front_matter(doc)

    text_section = doc.add_section(WD_SECTION.NEW_PAGE)
    text_section.page_width = Cm(21)
    text_section.page_height = Cm(29.7)
    text_section.top_margin = Cm(3)
    text_section.bottom_margin = Cm(2)
    text_section.left_margin = Cm(3)
    text_section.right_margin = Cm(2)
    add_page_number_section(text_section)

    add_introduction(doc)
    add_baseline(doc)
    add_evolution(doc, figures)
    add_architecture(doc, figures)
    add_functional_evolution(doc)
    add_security(doc, figures)
    add_use_cases_and_validation(doc, figures)
    add_limitations(doc)
    add_conclusion(doc)
    add_references(doc)
    add_appendix(doc)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build()
